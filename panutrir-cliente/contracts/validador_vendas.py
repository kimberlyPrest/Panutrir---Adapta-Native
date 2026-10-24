#!/usr/bin/env python3
"""Validador do contrato de vendas v2.0.0 e da agregacao de venda_mensal.

T2.1c / SPEC-1-002. Nao fala com banco: recebe .csv (o que a carga consome) ou
.xlsx/.xlsm (a congelada como vem do BI) e prova as regras do contrato. A carga
real e a T2.2.

Regras que este arquivo garante (e o motivo de cada uma):
  * SOMA TODOS OS TIPOS. Filtrar tipo='Venda' descarta as trocas e devolve venda
    liquida inflada — R$ 7,9 M em 2026. E a regra de ouro do contrato.
  * NEGATIVO NAO E ERRO. Os arquivos reais trazem estornos (~0,2% das linhas):
    vlritens negativo com qtde negativa em 'Venda', vlrtroca negativo em 'Troca'.
  * REJEICAO NAO INVALIDA O ARQUIVO. So e invalido o que impede a leitura
    (cabecalho sem coluna obrigatoria, arquivo vazio). Linha rejeitada e contada
    e listada — nada em silencio (AGENTS.md).
  * ACEITA 20 E 22 COLUNAS. 2024/2025 nao tem codcli_sap/codpro_sap.

Uso:
    python3 contracts/validador_vendas.py <arquivo.csv|.xlsx>
    python3 contracts/validador_vendas.py <arquivo> --json
    python3 contracts/validador_vendas.py <arquivo> --mes 2024-10 2026-09

Saida: resumo da validacao + agregado mensal. Codigo de saida 0 (legivel) ou 1 (ilegivel).
"""
from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

CONTRATO = Path(__file__).with_name("vendas.v2.0.0.json")

# ---- cabecalho real -> nome canonico do contrato ---------------------------
ALIASES = {
    "mesano: mês": "mes",
    "mesano: mes": "mes",
    "id_empresa": "id_empresa",
    "tipo": "tipo",
    "regiao": "regiao",
    "regional": "regional",
    "cd_resp": "cd_resp",
    "codrede": "codrede",
    "rede": "rede",
    "cidade": "cidade",
    "uf": "uf",
    "codcli": "cod_cliente",
    "codcli_sap": "codcli_sap",
    "fantasia": "fantasia",
    "codpro": "cod_sku",
    "codpro_sap": "codpro_sap",
    "produto": "produto",
    "linha": "linha",
    "soma de qtde": "qtdade",
    "soma de vlritens": "vlritens",
    "soma de vlrdesc": "vlrdesc",
    "soma de qtdetroca": "qtdetroca",
    "soma de vlrtroca": "vlrtroca",
}

TIPOS_VALIDOS = {
    "Venda", "Troca", "Devoluções Diversas", "Bonificação",
    "Qualidade", "Falta de Produtos", "Divergências", "Fora do Padrão",
}

OBRIGATORIAS = [
    "mes", "id_empresa", "tipo", "uf", "cod_cliente", "cod_sku",
    "produto", "linha", "qtdade", "vlritens", "vlrdesc", "qtdetroca", "vlrtroca",
]

LINHAS_DE_RODAPE = {"totais", "total"}


# ---------------------------------------------------------------- conversores
def _dec(valor) -> Decimal:
    """Numero no padrao brasileiro ou ingles. Levanta ValueError se invalido."""
    if valor is None:
        raise ValueError("vazio")
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    txt = str(valor).strip()
    if txt == "":
        raise ValueError("vazio")
    if "," in txt and "." in txt:
        txt = txt.replace(".", "").replace(",", ".")
    elif "," in txt:
        txt = txt.replace(",", ".")
    try:
        return Decimal(txt)
    except InvalidOperation as exc:
        raise ValueError(f"nao numerico: {valor!r}") from exc


def _mes(valor) -> str:
    """Normaliza para YYYY-MM. Aceita data do Excel, YYYY-MM, YYYY-MM-DD e dd/mm/aaaa."""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m")
    txt = str(valor or "").strip()
    if not txt:
        raise ValueError("vazio")
    for formato in ("%Y-%m-%d", "%Y-%m", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(txt[:10], formato).strftime("%Y-%m")
        except ValueError:
            continue
    raise ValueError(f"data invalida: {valor!r}")


# ------------------------------------------------------------------- leitores
def _abrir_xlsx(caminho: Path):
    """Abre a congelada .xlsx. Devolve (iterador, cabecalho, rotulo) ou (None, erro, None)."""
    try:
        import openpyxl
    except ImportError:
        return None, "openpyxl nao instalado — use CSV ou instale openpyxl", None
    try:
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except Exception as exc:  # arquivo corrompido ou formato diferente
        return None, f"nao foi possivel abrir o xlsx: {exc}", None
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    try:
        cab = next(it)
    except StopIteration:
        wb.close()
        return None, "planilha vazia", None
    cabecalho = [str(c).strip() if c is not None else "" for c in cab]

    def gerador():
        try:
            for row in it:
                yield ["" if v is None else v for v in row]
        finally:
            wb.close()

    return gerador(), cabecalho, "(xlsx)"


def _abrir_csv(caminho: Path):
    """Devolve (fh, leitor, cabecalho, delimitador) ou levanta ValueError."""
    fh = caminho.open(newline="", encoding="utf-8-sig")
    amostra = fh.read(4096)
    fh.seek(0)
    primeira = amostra.split("\n")[0] if amostra else ""
    delim = ";" if primeira.count(";") > primeira.count(",") else ","
    leitor = csv.reader(fh, delimiter=delim)
    try:
        cabecalho = next(leitor)
    except StopIteration:
        fh.close()
        raise ValueError("arquivo vazio")
    return fh, leitor, cabecalho, delim


# ------------------------------------------------------------------ principal
def validar(caminho: str, limite_meses=None) -> dict:
    """Valida o arquivo e devolve o agregado. Nao grava nada."""
    caminho_p = Path(caminho)
    if not caminho_p.exists():
        return {"ok": False, "erro": f"arquivo nao encontrado: {caminho}"}

    fh = None
    sufixo = caminho_p.suffix.lower()
    if sufixo in (".xlsx", ".xlsm"):
        leitor, cabecalho, delim = _abrir_xlsx(caminho_p)
        if leitor is None:
            return {"ok": False, "erro": cabecalho}
    else:
        try:
            fh, leitor, cabecalho, delim = _abrir_csv(caminho_p)
        except ValueError as exc:
            return {"ok": False, "erro": str(exc)}

    try:
        # mapeia colunas reais -> canonicas
        mapa = {}
        for i, nome in enumerate(cabecalho):
            canonico = ALIASES.get(str(nome).strip().lower())
            if canonico:
                mapa[i] = canonico

        presentes = set(mapa.values())
        faltando = [c for c in OBRIGATORIAS if c not in presentes]
        if faltando:
            return {
                "ok": False,
                "erro": "colunas obrigatorias ausentes",
                "faltando": faltando,
                "cabecalho_lido": cabecalho,
                "delimitador": delim,
            }

        agg = defaultdict(
            lambda: {"vlritens": Decimal("0"), "vlrdesc": Decimal("0"),
                     "vlrtroca": Decimal("0"), "qtdade": Decimal("0"), "linhas": 0}
        )
        rejeicoes = defaultdict(int)
        exemplos_rejeicao = []
        tipos_vistos = defaultdict(int)
        linhas_lidas = 0
        linhas_rodape = 0
        linhas_sem_valor = 0

        def rejeitar(motivo, linha_n, detalhe):
            rejeicoes[motivo] += 1
            if len(exemplos_rejeicao) < 10:
                exemplos_rejeicao.append(
                    {"linha": linha_n, "motivo": motivo, "detalhe": detalhe}
                )

        for linha_n, bruto in enumerate(leitor, start=2):
            if not bruto or all(str(c).strip() == "" for c in bruto):
                continue
            if str(bruto[0]).strip().lower() in LINHAS_DE_RODAPE:
                linhas_rodape += 1
                continue

            linhas_lidas += 1
            reg = {}
            for i, valor in enumerate(bruto):
                canonico = mapa.get(i)
                if canonico:
                    reg[canonico] = valor

            try:
                mes = _mes(reg.get("mes"))
            except ValueError as exc:
                rejeitar("data_invalida", linha_n, str(exc))
                continue

            if limite_meses and not (limite_meses[0] <= mes <= limite_meses[1]):
                rejeitar("fora_da_janela", linha_n, mes)
                continue

            try:
                cli = int(_dec(reg.get("cod_cliente")))
                sku = int(_dec(reg.get("cod_sku")))
            except (ValueError, TypeError) as exc:
                rejeitar("codigo_vazio", linha_n, str(exc))
                continue
            if cli <= 0 or sku <= 0:
                rejeitar("codigo_vazio", linha_n, "cli={} sku={}".format(cli, sku))
                continue

            tipo = str(reg.get("tipo", "")).strip()
            if tipo not in TIPOS_VALIDOS:
                rejeitar("tipo_desconhecido", linha_n, tipo or "(vazio)")
                continue
            tipos_vistos[tipo] += 1

            try:
                vb = _dec(reg.get("vlritens") or "0")
                vd = _dec(reg.get("vlrdesc") or "0")
                vt = _dec(reg.get("vlrtroca") or "0")
                qt = _dec(reg.get("qtdade") or "0")
            except ValueError as exc:
                rejeitar("valor_nao_numerico", linha_n, str(exc))
                continue
            # Negativo NAO e erro: os arquivos reais trazem estornos (~0,2%).
            # Somar como esta; truncar em zero inflaria a venda liquida.

            if vb == 0 and vd == 0 and vt == 0:
                linhas_sem_valor += 1
                continue

            a = agg[(mes, cli, sku)]
            a["vlritens"] += vb
            a["vlrdesc"] += vd
            a["vlrtroca"] += vt
            a["qtdade"] += qt
            a["linhas"] += 1
    finally:
        if fh is not None:
            fh.close()

    # ------------------------------------------------------------- totais
    tot = {"vlritens": Decimal("0"), "vlrdesc": Decimal("0"),
           "vlrtroca": Decimal("0"), "qtdade": Decimal("0")}
    por_mes = defaultdict(
        lambda: {"vlritens": Decimal("0"), "vlrdesc": Decimal("0"),
                 "vlrtroca": Decimal("0"), "chaves": set()}
    )
    for (mes, cli, sku), a in agg.items():
        for k in ("vlritens", "vlrdesc", "vlrtroca", "qtdade"):
            tot[k] += a[k]
        pm = por_mes[mes]
        for k in ("vlritens", "vlrdesc", "vlrtroca"):
            pm[k] += a[k]
        pm["chaves"].add((cli, sku))

    return {
        # V-07: rejeicao nao invalida o arquivo; so e invalido o ilegivel.
        "ok": True,
        "arquivo": caminho,
        "delimitador": delim,
        "linhas_lidas": linhas_lidas,
        "linhas_de_rodape_ignoradas": linhas_rodape,
        "linhas_sem_valor_ignoradas": linhas_sem_valor,
        "linhas_rejeitadas": sum(rejeicoes.values()),
        "combos_na_saida": len(agg),
        "tipos_vistos": dict(tipos_vistos),
        "rejeicoes_por_motivo": dict(rejeicoes),
        "exemplos_rejeicao": exemplos_rejeicao,
        "totais": {k: str(v) for k, v in tot.items()},
        "venda_liquida": str(tot["vlritens"] - tot["vlrdesc"] - tot["vlrtroca"]),
        "por_mes": {
            m: {
                "vlritens": str(d["vlritens"]),
                "vlrdesc": str(d["vlrdesc"]),
                "vlrtroca": str(d["vlrtroca"]),
                "venda_liquida": str(d["vlritens"] - d["vlrdesc"] - d["vlrtroca"]),
                "combos": len(d["chaves"]),
            }
            for m, d in sorted(por_mes.items())
        },
        # saida pronta para a carga (T2.2 consome isto)
        "venda_mensal": [
            {
                "mes": mes,
                "cod_cliente": cli,
                "cod_sku": sku,
                "vlritens": str(a["vlritens"]),
                "vlrdesc": str(a["vlrdesc"]),
                "vlrtroca": str(a["vlrtroca"]),
                "venda_bruta": str(a["vlritens"]),
                "venda_liquida": str(a["vlritens"] - a["vlrdesc"] - a["vlrtroca"]),
                "qtdade": str(a["qtdade"]),
            }
            for (mes, cli, sku), a in sorted(agg.items())
        ],
    }


def resumo(r: dict) -> str:
    if not r.get("ok"):
        linhas = ["INVALIDO: {}".format(r.get("erro"))]
        if r.get("faltando"):
            linhas.append("  colunas faltando: {}".format(r["faltando"]))
        return "\n".join(linhas)

    linhas = [
        "arquivo ....: {}".format(r["arquivo"]),
        "delimitador : {!r}".format(r["delimitador"]),
        "linhas lidas: {}".format(r["linhas_lidas"]),
        "  rodape ignorado .......: {}".format(r["linhas_de_rodape_ignoradas"]),
        "  sem valor (ignoradas) .: {}".format(r["linhas_sem_valor_ignoradas"]),
        "  rejeitadas ............: {}".format(r["linhas_rejeitadas"]),
        "combos (mes,cli,sku) .....: {}".format(r["combos_na_saida"]),
        "tipos vistos .............: {}".format(r["tipos_vistos"]),
        "vlritens .................: {}".format(r["totais"]["vlritens"]),
        "vlrdesc ..................: {}".format(r["totais"]["vlrdesc"]),
        "vlrtroca .................: {}".format(r["totais"]["vlrtroca"]),
        "venda liquida ............: {}".format(r["venda_liquida"]),
    ]
    if r["rejeicoes_por_motivo"]:
        linhas.append("REJEICOES: {}".format(r["rejeicoes_por_motivo"]))
        for e in r["exemplos_rejeicao"]:
            linhas.append("  linha {}: {} — {}".format(e["linha"], e["motivo"], e["detalhe"]))
    linhas.append("")
    linhas.append("{:8} {:>7} {:>16} {:>12} {:>14} {:>16}".format(
        "mes", "combos", "vlritens", "vlrdesc", "vlrtroca", "liquida"))
    for m, d in r["por_mes"].items():
        linhas.append("{:8} {:>7} {:>16} {:>12} {:>14} {:>16}".format(
            m, d["combos"], d["vlritens"], d["vlrdesc"], d["vlrtroca"], d["venda_liquida"]))
    linhas.append("")
    if r["linhas_rejeitadas"]:
        linhas.append(
            "RESULTADO: VALIDO, com {} linha(s) rejeitada(s) listada(s) acima "
            "(nao bloqueiam a carga)".format(r["linhas_rejeitadas"])
        )
    else:
        linhas.append("RESULTADO: VALIDO — nenhuma rejeicao")
    return "\n".join(linhas)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    limite = None
    if "--mes" in sys.argv:
        i = sys.argv.index("--mes")
        limite = (sys.argv[i + 1], sys.argv[i + 2])
    res = validar(sys.argv[1], limite)
    if "--json" in sys.argv:
        print(json.dumps({k: v for k, v in res.items() if k != "venda_mensal"},
                         ensure_ascii=False, indent=2))
    else:
        print(resumo(res))
    sys.exit(0 if res.get("ok") else 1)
